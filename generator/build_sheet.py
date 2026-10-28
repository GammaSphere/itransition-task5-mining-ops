"""Builds the generator workbook (generator/WY_Mining_Ops_Simulator.xlsx).

Upload it to Google Drive with conversion to Google Sheets. All logic lives in
cell formulas (Google Sheets array formulas: LET, LAMBDA, MAP, SCAN, MAKEARRAY,
BYROW, FILTER); there is no script. Formulas mirror generator/model.py.

Sheets:
  Generator  parameters + live charts (the only sheet a user edits)
  Data       the data source: Date | one column per mine | Total (one spill formula)
  Engine     day-by-day calculation, one array formula per column
  Model      derived distribution constants and the event -> mine map
  Guide      how the model works
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from model import MAX_EVENTS, MAX_MINES, Params

OUT = Path(__file__).with_name("WY_Mining_Ops_Simulator.xlsx")
MAX_DAYS = 730
LAST_ROW = MAX_DAYS + 2  # Engine data starts on row 3

# Reference categorical palette (fixed order; colour follows the mine slot).
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948"]
INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
LABEL_BG, HEAD_BG, BAND_BG, CALC_BG, NOTE_BG = "D9D9D9", "EFEFEF", "BFBFBF", "F6F6F4", "FFF8E1"

thin = Side(style="thin", color="B7B7B7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

HASH = "MOD(SIN(Seed*0.6180339887+{t}*12.9898+{k}*78.233+{s}*37.719)*43758.5453,1)"


def gauss(t: str, k: str, s: str) -> str:
    """Standard normal draw from the deterministic hash (clamped away from 0 and 1)."""
    return f"NORMSINV(MIN(MAX({HASH.format(t=t, k=k, s=s)},1E-9),1-1E-9))"


def rng(sheet: str, c: str) -> str:
    """Engine column c, rows 3 .. DaysCount+2 (exactly one row per simulated day)."""
    p = f"{sheet}!" if sheet else ""
    return f"{p}{c}3:INDEX({p}{c}:{c},DaysCount+2)"


# Engine column layout -------------------------------------------------------
E_T, E_DATE, E_WD, E_TREND, E_COMMON = "A", "B", "C", "D", "E"
Z0, BASE0, W0, EVF0, OUT0 = 6, 6 + MAX_MINES, 6 + 2 * MAX_MINES, 6 + 2 * MAX_MINES + MAX_EVENTS, 6 + 3 * MAX_MINES + MAX_EVENTS
E_TOTAL = col(OUT0 + MAX_MINES)


def zc(m): return col(Z0 + m - 1)
def bc(m): return col(BASE0 + m - 1)
def wc(e): return col(W0 + e - 1)
def fc(m): return col(EVF0 + m - 1)
def oc(m): return col(OUT0 + m - 1)


def style_range(ws, ref, *, font=None, fill=None, align=None, border=None, fmt=None):
    for row in ws[ref]:
        for cell in row:
            if font: cell.font = font
            if fill: cell.fill = PatternFill("solid", fgColor=fill)
            if align: cell.alignment = align
            if border: cell.border = border
            if fmt: cell.number_format = fmt


def name(wb, key, ref):
    wb.defined_names[key] = DefinedName(key, attr_text=ref)


HINTS = {
    4: f"Start date and number of days (1 to {MAX_DAYS}). Grey italic cells are calculated.",
    5: f"Up to {MAX_MINES} mines: type a name to add a mine, clear it to remove one.",
    6: "The parameter headers follow the distribution: Normal / Log-normal: Mean, Deviation; "
       "Uniform: Min, Max; Triangular: Min, Mode, Max.",
    7: "Smoothing: day-to-day correlation, 0 to 0.99 (0 = white noise). Correlation: between mines, 0 to 1. "
       "Seed: change it for a new random realisation.",
    10: "Weekday multipliers: 1 = normal output, 0.4 = output reduced by 60 %.",
    11: "Trend per day. Compound: x(1+g)^t, Linear: x(1+g*t). Negative = decline.",
    12: "Day = first day of the event. Factor at the peak: 0.4 = drop to 40 %, 1.4 = spike +40 %.",
    13: "Duration in days; for Bell it is the width at half height (smooth ramp up and fade out).",
    14: "Probability: chance that each mine is hit (independent per mine, blank = 1).",
    15: "Mines: exact number of mines hit (optional, overrides Probability).",
    16: "Shape: Bell (bell-shaped curve) or Flat (rectangular step).",
}


CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build_generator(wb, p: Params):
    ws = wb.active
    ws.title = "Generator"
    ws.sheet_properties.tabColor = "2A78D6"
    ws.sheet_view.showGridLines = False
    widths = {"A": 16, "B": 14, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for r in range(1, 23):
        ws.row_dimensions[r].height = 22
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[8].height = 30
    ws.row_dimensions[21].height = 62

    ws.merge_cells("A1:I1")
    ws["A1"] = "Weyland-Yutani Corporation"
    ws["A1"].font = Font(size=20, color=INK)
    ws["A1"].alignment = CENTER
    ws.merge_cells("A2:I2")
    ws["A2"] = "Mining Ops Simulator"
    ws["A2"].font = Font(bold=True, size=12, color=INK)
    style_range(ws, "A2:I2", fill=BAND_BG, align=CENTER, border=BOX)

    # Grid frame: labels column A, cells B..I
    style_range(ws, "A3:I20", border=BOX, align=CENTER)
    style_range(ws, "A3:A20", fill=LABEL_BG, font=Font(bold=True, color=INK), align=RIGHT)
    for r, label in {3: "Range:", 5: "Mines:", 6: "Random:", 9: "Weekday:", 10: "Daily output:",
                     11: "Daily growth:", 12: "Events:"}.items():
        ws.cell(row=r, column=1, value=label)

    head = Font(bold=True, color=INK)
    calc = Font(italic=True, color=INK2)

    # Range + live summary -----------------------------------------------------
    for c, h in zip("BCDEFGHI", ["Start date", "Days count", "End date", "Mines", "Total output",
                                 "Avg / day", "Peak day", "Peak total"]):
        ws[f"{c}3"] = h
        ws[f"{c}3"].font = head
    ws["B4"] = p.start
    ws["B4"].number_format = "yyyy-mm-dd"
    ws["C4"] = p.days
    ws["D4"] = '=IF(N(C4)<1,"",B4+C4-1)'
    ws["E4"] = "=COUNTA(MineNames)"
    # Bounded ranges only: the xlsx importer parses formulas with Excel grammar, where "AT3:AT" is invalid.
    total = rng("Engine", E_TOTAL)
    ws["F4"] = f"=IFERROR(SUM({total}),0)"
    ws["G4"] = "=IFERROR(F4/C4,0)"
    ws["H4"] = f'=IFERROR(INDEX({rng("Engine", E_DATE)},MATCH(MAX({total}),{total},0)),"")'
    ws["I4"] = f'=IFERROR(MAX({total}),"")'
    for ref, fmt in {"D4": "yyyy-mm-dd", "E4": "0", "F4": "#,##0", "G4": "#,##0.0", "H4": "yyyy-mm-dd", "I4": "#,##0.0"}.items():
        ws[ref].number_format = fmt
        ws[ref].font = calc
        ws[ref].fill = PatternFill("solid", fgColor=CALC_BG)

    # Mines ---------------------------------------------------------------------
    for i in range(MAX_MINES):
        cell = ws.cell(row=5, column=2 + i, value=p.mines[i] if i < len(p.mines) else None)
        cell.font = Font(color=INK)

    # Random --------------------------------------------------------------------
    ws["B6"] = "Distribution"
    ws["B6"].font = head
    ws["C6"] = '=SWITCH(Dist,"Normal","Mean","Uniform","Min","Log-normal","Mean","Triangular","Min","Param 1")'
    ws["D6"] = '=SWITCH(Dist,"Normal","Deviation","Uniform","Max","Log-normal","Deviation","Triangular","Mode","Param 2")'
    ws["E6"] = '=IF(Dist="Triangular","Max","")'
    for ref in ("C6", "D6", "E6"):
        ws[ref].font = head
    ws["F6"], ws["G6"], ws["H6"] = "Smoothing", "Correlation", "Seed"
    for ref in ("F6", "G6", "H6"):
        ws[ref].font = head
    ws["B7"] = p.distribution
    ws["C7"], ws["D7"], ws["E7"] = p.a, p.b, p.c
    ws["F7"], ws["G7"], ws["H7"] = p.smoothing, p.correlation, p.seed
    ws.merge_cells("B8:I8")
    ws["B8"] = (
        '=SWITCH(Dist,'
        '"Normal","Normal: mean "&TEXT(NormMu,"0.0")&", deviation "&TEXT(NormSd,"0.0"),'
        '"Uniform","Uniform: every value between "&TEXT(UniLo,"0.0")&" and "&TEXT(UniHi,"0.0")&" is equally likely",'
        '"Log-normal","Log-normal (right-skewed, never negative): log-mean "&TEXT(LogMu,"0.000")&", log-sd "&TEXT(LogSd,"0.000"),'
        '"Triangular","Triangular: min "&TEXT(TriLo,"0.0")&", most likely "&TEXT(TriMode,"0.0")&", max "&TEXT(TriHi,"0.0"),'
        '"Choose a distribution")&"   |   base mean "&TEXT(BaseMean,"0.0")&", sd "&TEXT(BaseSd,"0.0")'
        '&"   |   day-to-day corr. "&TEXT(Rho,"0.00")&", between mines "&TEXT(MixC,"0.00")'
    )
    ws["B8"].font = Font(italic=True, size=9, color=INK2)
    ws["B8"].alignment = LEFT
    ws["A8"].fill = PatternFill("solid", fgColor=LABEL_BG)

    # Day-of-week factors -------------------------------------------------------
    for i, d in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]):
        ws.cell(row=9, column=2 + i, value=d).font = head
        ws.cell(row=10, column=2 + i, value=p.weekday[i])

    # Trend ---------------------------------------------------------------------
    ws["B11"] = p.growth
    ws["B11"].number_format = "0.0%"
    ws["C11"] = "Trend"
    ws["C11"].font = head
    ws["D11"] = p.trend

    # Events --------------------------------------------------------------------
    for c, h in zip("BCDEFG", ["Day", "Duration", "Factor", "Probability", "Mines", "Shape"]):
        ws[f"{c}12"] = h
        ws[f"{c}12"].font = head
    ws.merge_cells("H12:I12")
    ws["H12"] = "Affected mines"
    ws["H12"].font = head
    for e in range(MAX_EVENTS):
        r = 13 + e
        ws.cell(row=r, column=1, value=f"#{e + 1}").font = Font(color=MUTED)
        ev = p.events[e] if e < len(p.events) else None
        if ev:
            ws.cell(row=r, column=2, value=ev.day)
            ws.cell(row=r, column=3, value=ev.duration)
            ws.cell(row=r, column=4, value=ev.factor)
            ws.cell(row=r, column=5, value=ev.probability)
            ws.cell(row=r, column=6, value=ev.mines)
            ws.cell(row=r, column=7, value=ev.shape)
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.merge_cells(f"H{r}:I{r}")
        ws[f"H{r}"] = (f'=IF(OR(B{r}="",D{r}=""),"",IFERROR(TEXTJOIN(", ",TRUE,'
                       f'FILTER(MineNames,INDEX(EvHit,{e + 1},0)=1)),"none"))')
        ws[f"H{r}"].font = calc
        style_range(ws, f"H{r}:I{r}", fill=CALC_BG)

    ws.column_dimensions["J"].width = 70
    for r, text in HINTS.items():
        ws.cell(row=r, column=10, value=text).font = Font(italic=True, size=9, color=MUTED)

    ws.merge_cells("A21:I21")
    ws["A21"] = (
        "How it works: each mine draws daily base output from the chosen distribution (Gaussian-copula with "
        "day-to-day smoothing and between-mine correlation), then multiplies it by the weekday factor, the trend and "
        "every event that hits it. The Data sheet is the data source for the dashboard. Details: Guide sheet."
    )
    ws["A21"].font = Font(size=9, color=INK2)
    ws["A21"].alignment = LEFT
    ws["A21"].fill = PatternFill("solid", fgColor=NOTE_BG)

    # Conditional formatting: third parameter only matters for Triangular
    ws.conditional_formatting.add("E7", FormulaRule(formula=['$B$7<>"Triangular"'], fill=PatternFill("solid", fgColor="EEEEEE"), font=Font(color="BBBBBB")))
    ws.conditional_formatting.add("F13:F20", FormulaRule(formula=['AND($F13<>"",$B13<>"")'], font=Font(bold=True, color=INK)))
    ws.conditional_formatting.add("E13:E20", FormulaRule(formula=['AND($F13<>"",$B13<>"")'], font=Font(strike=True, color="AAAAAA")))
    ws.conditional_formatting.add("D13:D20", FormulaRule(formula=['AND(ISNUMBER($D13),$D13<1)'], font=Font(color="C0392B", bold=True)))
    ws.conditional_formatting.add("D13:D20", FormulaRule(formula=['AND(ISNUMBER($D13),$D13>1)'], font=Font(color="1E7B34", bold=True)))

    # Validation ------------------------------------------------------------------
    def dv(ref, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
        ws.add_data_validation(v)
        v.add(ref)

    dv("B4", type="date", operator="greaterThan", formula1="1", error="Enter a date")
    dv("C4", type="whole", operator="between", formula1="1", formula2=str(MAX_DAYS), error=f"1 to {MAX_DAYS} days")
    dv("B7", type="list", formula1='"Normal,Uniform,Log-normal,Triangular"')
    dv("F7", type="decimal", operator="between", formula1="0", formula2="0.99", error="Smoothing must be 0 to 0.99")
    dv("G7", type="decimal", operator="between", formula1="0", formula2="1", error="Correlation must be 0 to 1")
    dv("B10:H10", type="decimal", operator="greaterThanOrEqual", formula1="0", error="Factor must be >= 0")
    dv("D11", type="list", formula1='"Compound,Linear"')
    dv("B13:B20", type="date", operator="greaterThan", formula1="1", error="Enter a date")
    dv("C13:C20", type="decimal", operator="between", formula1="1", formula2="365", error="Duration 1 to 365 days")
    dv("D13:D20", type="decimal", operator="greaterThanOrEqual", formula1="0", error="Factor must be >= 0")
    dv("E13:E20", type="decimal", operator="between", formula1="0", formula2="1", error="Probability 0 to 1")
    dv("F13:F20", type="whole", operator="between", formula1="0", formula2=str(MAX_MINES), error=f"0 to {MAX_MINES} mines")
    dv("G13:G20", type="list", formula1='"Bell,Flat"')

    # Named inputs ----------------------------------------------------------------
    for key, ref in {
        "StartDate": "$B$4", "DaysCount": "$C$4", "MineNames": f"$B$5:${col(1 + MAX_MINES)}$5",
        "Dist": "$B$7", "ParamA": "$C$7", "ParamB": "$D$7", "ParamC": "$E$7",
        "Smoothing": "$F$7", "Correlation": "$G$7", "Seed": "$H$7",
        "WeekdayFactors": "$B$10:$H$10", "GrowthRate": "$B$11", "TrendType": "$D$11",
        "EvDay": "$B$13:$B$20", "EvDur": "$C$13:$C$20", "EvFactor": "$D$13:$D$20",
        "EvProb": "$E$13:$E$20", "EvMines": "$F$13:$F$20", "EvShape": "$G$13:$G$20",
    }.items():
        name(wb, key, f"Generator!{ref}")
    return ws


def build_engine(wb):
    ws = wb.create_sheet("Engine")
    ws.sheet_properties.tabColor = "898781"
    ws.freeze_panes = "C3"
    groups = [
        (E_T, E_COMMON, "Calendar, weekday & trend"),
        (zc(1), zc(MAX_MINES), "Noise z (standard normal, smoothed & correlated)"),
        (bc(1), bc(MAX_MINES), "Base output = F^-1(Phi(z)) of chosen distribution"),
        (wc(1), wc(MAX_EVENTS), "Event shape weight w(t) (0..1)"),
        (fc(1), fc(MAX_MINES), "Event factor = PRODUCT(1 + (F-1) w hit)"),
        (oc(1), E_TOTAL, "Output = max(0, base) x weekday x trend x events"),
    ]
    fills = ["EFEFEF", "E8F1FB", "EAF6F0", "FDF3E1", "FBEDEA", "EDEBFA"]
    for (a, b, title), fill in zip(groups, fills):
        if a != b:
            ws.merge_cells(f"{a}1:{b}1")
        ws[f"{a}1"] = title
        style_range(ws, f"{a}1:{b}2", fill=fill, font=Font(bold=True, size=9, color=INK), align=CENTER, border=BOX)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30

    ws["A2"], ws["B2"], ws["C2"], ws["D2"], ws["E2"] = "t", "Date", "Weekday x", "Trend x", "Common shock"
    ws["A3"] = "=SEQUENCE(DaysCount)"
    ws["B3"] = "=SEQUENCE(DaysCount,1,StartDate,1)"
    ws["C3"] = f"=MAP({rng('', E_DATE)},LAMBDA(d,INDEX(WeekdayFactors,1,WEEKDAY(d,2))))"
    ws["D3"] = (f'=MAP({rng("", E_T)},LAMBDA(t,IF(TrendType="Linear",MAX(0,1+GrowthRate*(t-1)),(1+GrowthRate)^(t-1))))')
    ws["E3"] = f"=MAP({rng('', E_T)},LAMBDA(t,{gauss('t', '0', '1')}))"

    slot_name = lambda m: f'INDEX(MineNames,1,{m})'
    for m in range(1, MAX_MINES + 1):
        label = f'=IF({slot_name(m)}="","slot {m}",{slot_name(m)})'
        ws[f"{zc(m)}2"] = label
        ws[f"{bc(m)}2"] = label
        ws[f"{fc(m)}2"] = label
        ws[f"{oc(m)}2"] = f'=IF({slot_name(m)}="","",{slot_name(m)})'
        # AR(1) in a Gaussian copula: z0 ~ N(0,1); z_t = rho z_(t-1) + sqrt(1-rho^2) eps_t
        eps = (f"MAP({rng('', E_T)},{rng('', E_COMMON)},LAMBDA(t,cs,"
               f"SQRT(MixC)*cs+SQRT(1-MixC)*{gauss('t', str(m), '1')}))")
        ws[f"{zc(m)}3"] = f"=SCAN({gauss('0', str(m), '2')},{eps},LAMBDA(z,e,Rho*z+SQRT(1-Rho^2)*e))"
        ws[f"{bc(m)}3"] = (
            f"=MAP({rng('', zc(m))},LAMBDA(z,SWITCH(Dist,"
            f'"Uniform",UniLo+(UniHi-UniLo)*NORMSDIST(z),'
            f'"Log-normal",EXP(LogMu+LogSd*z),'
            f'"Triangular",IF(TriHi=TriLo,TriLo,LET(q,NORMSDIST(z),f,(TriMode-TriLo)/(TriHi-TriLo),'
            f"IF(q<f,TriLo+SQRT(q*(TriHi-TriLo)*(TriMode-TriLo)),TriHi-SQRT((1-q)*(TriHi-TriLo)*(TriHi-TriMode))))),"
            f"NormMu+NormSd*z)))"
        )
        ws[f"{fc(m)}3"] = (f"=LET(amp,TRANSPOSE(INDEX(EvAmp,0,{m})),"
                           f"BYROW({wc(1)}3:INDEX({wc(MAX_EVENTS)}:{wc(MAX_EVENTS)},DaysCount+2),"
                           f"LAMBDA(w,PRODUCT(ARRAYFORMULA(1+amp*w)))))")
        ws[f"{oc(m)}3"] = (f"=MAP({rng('', bc(m))},{rng('', E_WD)},{rng('', E_TREND)},{rng('', fc(m))},"
                           f'LAMBDA(b,wd,tr,ev,IF({slot_name(m)}="",,ROUND(MAX(0,b)*wd*tr*ev,2))))')
    for e in range(1, MAX_EVENTS + 1):
        ws[f"{wc(e)}2"] = f"Event {e}"
        ws[f"{wc(e)}3"] = (
            f"=LET(dstart,INDEX(EvDay,{e}),fac,INDEX(EvFactor,{e}),dur,MAX(1,N(INDEX(EvDur,{e}))),"
            f'flat,INDEX(EvShape,{e})="Flat",ctr,dstart+(dur-1)/2,sg,dur/2.3548,'
            f"MAP({rng('', E_DATE)},LAMBDA(d,IF(OR(dstart=\"\",fac=\"\"),0,"
            f"IF(flat,IF(AND(d>=dstart,d<dstart+dur),1,0),"
            f"LET(w,EXP(-((d-ctr)^2)/(2*sg^2)),IF(w<0.01,0,w)))))))"
        )
    ws[f"{E_TOTAL}2"] = "Total"
    ws[f"{E_TOTAL}3"] = f"=BYROW({oc(1)}3:INDEX({oc(MAX_MINES)}:{oc(MAX_MINES)},DaysCount+2),LAMBDA(r,SUM(r)))"

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 11
    for c in range(3, OUT0 + MAX_MINES + 1):
        ws.column_dimensions[col(c)].width = 10
    ws.column_dimensions["B"].number_format = "yyyy-mm-dd"
    return ws


def build_model(wb):
    ws = wb.create_sheet("Model")
    ws.sheet_properties.tabColor = "898781"
    ws.column_dimensions["A"].width = 30
    for c in range(2, 2 + MAX_MINES):
        ws.column_dimensions[col(c)].width = 12
    ws.column_dimensions[col(2 + MAX_MINES)].width = 60
    ws["A1"] = "Model: derived constants and event map (calculated, do not edit)"
    ws["A1"].font = Font(bold=True, size=12)
    consts = [
        ("Distribution", "=Dist", None, ""),
        ("Normal mean", "=N(ParamA)", "NormMu", "Mean"),
        ("Normal deviation", "=MAX(0,N(ParamB))", "NormSd", "Deviation, negative treated as 0"),
        ("Uniform min", "=MIN(N(ParamA),N(ParamB))", "UniLo", "Min/Max are sorted, so a swapped pair still works"),
        ("Uniform max", "=MAX(N(ParamA),N(ParamB))", "UniHi", ""),
        ("Log-normal log-sd", "=SQRT(LN(1+(MAX(0,N(ParamB))/MAX(N(ParamA),1E-9))^2))", "LogSd", "sigma = sqrt(ln(1 + (sd/mean)^2))"),
        ("Log-normal log-mean", "=LN(MAX(N(ParamA),1E-9))-B8^2/2", "LogMu", "mu = ln(mean) - sigma^2/2"),
        ("Triangular min", "=MIN(ParamA,ParamB,ParamC)", "TriLo", "Min, Mode, Max are sorted"),
        ("Triangular mode", "=MEDIAN(ParamA,ParamB,ParamC)", "TriMode", ""),
        ("Triangular max", "=MAX(ParamA,ParamB,ParamC)", "TriHi", ""),
        ("Base mean", '=SWITCH(Dist,"Uniform",(B6+B7)/2,"Log-normal",MAX(N(ParamA),0),"Triangular",(B10+B11+B12)/3,B4)', "BaseMean", "Mean of the base distribution"),
        ("Base deviation", '=SWITCH(Dist,"Uniform",(B7-B6)/SQRT(12),"Log-normal",MAX(0,N(ParamB)),"Triangular",SQRT((B10^2+B11^2+B12^2-B10*B11-B10*B12-B11*B12)/18),B5)', "BaseSd", ""),
        ("Smoothing rho (clamped)", "=MIN(MAX(N(Smoothing),0),0.99)", "Rho", "Day-to-day autocorrelation"),
        ("Mine correlation c (clamped)", "=MIN(MAX(N(Correlation),0),1)", "MixC", "Share of the common daily shock"),
    ]
    for i, (label, f, key, hint) in enumerate(consts, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=f).number_format = "0.0000"
        ws.cell(row=i, column=3, value=hint).font = Font(italic=True, color=INK2)
        if key:
            name(wb, key, f"Model!$B${i}")

    last_c = col(1 + MAX_MINES)

    def block(top, title, formula, key, fmt):
        ws.cell(row=top, column=1, value=title).font = Font(bold=True)
        for m in range(1, MAX_MINES + 1):
            ws.cell(row=top, column=1 + m, value=f'=IF(INDEX(MineNames,1,{m})="","slot {m}",INDEX(MineNames,1,{m}))').font = Font(bold=True, color=INK2)
        for e in range(1, MAX_EVENTS + 1):
            ws.cell(row=top + e, column=1, value=f"Event {e}")
        ws.cell(row=top + 1, column=2, value=formula)
        style_range(ws, f"B{top + 1}:{last_c}{top + MAX_EVENTS}", fmt=fmt)
        name(wb, key, f"Model!$B${top + 1}:${last_c}${top + MAX_EVENTS}")
        return top + MAX_EVENTS + 2

    top = 19
    ws.cell(row=top - 1, column=1, value="Which mines does each event hit?").font = Font(bold=True, size=11)
    draw = f"=MAKEARRAY({MAX_EVENTS},{MAX_MINES},LAMBDA(e,m,{HASH.format(t='e', k='m', s='3')}))"
    top = block(top, "Random draw r(e, mine)", draw, "EvDraw", "0.000")
    hit = (
        f"=MAKEARRAY({MAX_EVENTS},{MAX_MINES},LAMBDA(e,m,LET(r,INDEX(EvDraw,e,m),"
        f'on,AND(INDEX(MineNames,1,m)<>"",INDEX(EvDay,e)<>"",INDEX(EvFactor,e)<>""),'
        f"cnt,INDEX(EvMines,e),pr,INDEX(EvProb,e),"
        f'IF(NOT(on),0,IF(cnt<>"",'
        f"IF(1+SUM(MAP(SEQUENCE(1,{MAX_MINES}),LAMBDA(k,IF(AND(INDEX(MineNames,1,k)<>\"\",INDEX(EvDraw,e,k)<r),1,0))))<=cnt,1,0),"
        f'IF(r<IF(pr="",1,pr),1,0))))))'
    )
    top = block(top, "Hit (1 = event affects the mine)", hit, "EvHit", "0")
    amp = f"=MAKEARRAY({MAX_EVENTS},{MAX_MINES},LAMBDA(e,m,INDEX(EvHit,e,m)*(N(INDEX(EvFactor,e))-1)))"
    block(top, "Amplitude (Factor - 1) x hit", amp, "EvAmp", "0.00")
    ws.cell(row=19, column=2 + MAX_MINES, value=(
        "Probability mode: a mine is hit when its draw r < Probability (blank = 1). "
        "Mines mode: the N named mines with the smallest draws are hit."
    )).alignment = LEFT
    return ws


def build_data(wb):
    ws = wb.create_sheet("Data", 1)
    ws.sheet_properties.tabColor = "1BAF7A"
    ws.freeze_panes = "B2"
    ws["A1"] = (
        '=IF(COUNTA(MineNames)=0,"Enter at least one mine name on the Generator sheet",'
        "LET(keep,ARRAYFORMULA(MineNames<>\"\"),"
        'VSTACK(HSTACK("Date",FILTER(MineNames,keep),"Total"),'
        f"HSTACK({rng('Engine', E_DATE)},"
        f"FILTER(Engine!{oc(1)}3:INDEX(Engine!{oc(MAX_MINES)}:{oc(MAX_MINES)},DaysCount+2),keep),"
        f"{rng('Engine', E_TOTAL)}))))"
    )
    ws.column_dimensions["A"].width = 12
    for c in range(2, MAX_MINES + 3):
        ws.column_dimensions[col(c)].width = 13
    style_range(ws, f"A1:{col(MAX_MINES + 2)}1", font=Font(bold=True), fill=HEAD_BG)
    ws.column_dimensions["A"].number_format = "yyyy-mm-dd"
    return ws


def build_guide(wb):
    ws = wb.create_sheet("Guide")
    ws.sheet_properties.tabColor = "EDA100"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Mining Ops Simulator: guide", "h"),
        ("Edit only the white cells on the Generator sheet. Everything else is formulas; the chart and the Data sheet update instantly.", ""),
        ("", ""),
        ("Parameters", "h2"),
        ("Range: start date and number of days (1 to 730).", ""),
        ("Mines: up to 8 names in row 5. Add or clear a name to change the number of mines.", ""),
        ("Distribution: Normal (Mean, Deviation), Uniform (Min, Max), Log-normal (Mean, Deviation), Triangular (Min, Mode, Max). The parameter headers change with the choice.", ""),
        ("Smoothing: day-to-day autocorrelation rho. 0 = white noise; 0.8 = smooth, slowly drifting output.", ""),
        ("Correlation: how strongly the mines move together (share of a common daily shock).", ""),
        ("Seed: the random numbers are a deterministic hash of (seed, day, mine), so the data only changes when a parameter changes.", ""),
        ("Daily output: one multiplier per weekday (e.g. Sunday 0.4 = output reduced by 60 %).", ""),
        ("Daily growth: trend per day, compound or linear.", ""),
        ("Events: day, duration, factor (0.4 = drop to 40 %, 1.4 = +40 % spike), probability that each mine is hit, or an exact number of mines (Mines column, overrides probability), and the shape (Bell or Flat).", ""),
        ("", ""),
        ("Model", "h2"),
        ("1. Shocks: eps(m,t) = sqrt(c) * common(t) + sqrt(1-c) * own(m,t), with common and own standard normal draws.", ""),
        ("2. Smoothing (AR(1)): z(m,t) = rho * z(m,t-1) + sqrt(1-rho^2) * eps(m,t). z stays standard normal, so the marginal distribution is not distorted.", ""),
        ("3. Gaussian copula: base(m,t) = F^-1(Phi(z(m,t))) where F is the chosen distribution. Consecutive days are correlated, values follow the distribution exactly.", ""),
        ("4. Output(m,t) = max(0, base) x weekday factor x trend(t) x PRODUCT over events (1 + (Factor-1) x w(t) x hit(e,m)), rounded to 0.01.", ""),
        ("   Bell events: w(t) = exp(-(t - centre)^2 / (2 sigma^2)) with sigma = Duration / 2.3548 (full width at half maximum = Duration).", ""),
        ("", ""),
        ("Sheets", "h2"),
        ("Data: the data source (Date | one column per mine | Total), read by the dashboard through the Google Visualization API.", ""),
        ("Engine: one array formula per column (MAP / SCAN / BYROW), one row per day.", ""),
        ("Model: distribution constants and the event-to-mine map (MAKEARRAY).", ""),
    ]
    for i, (text, kind) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h":
            c.font = Font(bold=True, size=14)
        elif kind == "h2":
            c.font = Font(bold=True, size=11, color="2A78D6")


def add_charts(wb, gen):
    eng = wb["Engine"]
    dates = Reference(eng, min_col=2, min_row=3, max_row=LAST_ROW)

    lc = LineChart()
    lc.title = "Daily output by mine"
    lc.y_axis.title = "Output"
    lc.x_axis.number_format = "mmm d"
    lc.height, lc.width = 9.5, 24
    lc.legend.position = "b"
    for m in range(1, MAX_MINES + 1):
        ref = Reference(eng, min_col=OUT0 + m - 1, min_row=3, max_row=LAST_ROW)
        lc.add_data(ref, titles_from_data=False)
        s = lc.series[-1]
        s.tx = SeriesLabel(strRef=StrRef(f"'Engine'!${oc(m)}$2"))
        s.smooth = False
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = SERIES[m - 1]
        s.graphicalProperties.line.width = 19050
    lc.set_categories(dates)
    lc.x_axis.delete = False
    lc.y_axis.delete = False
    gen.add_chart(lc, "A23")

    bc_ = BarChart()
    bc_.type = "col"
    bc_.title = "Total daily output (all mines)"
    bc_.height, bc_.width = 7.5, 24
    bc_.gapWidth = 40
    bc_.add_data(Reference(eng, min_col=OUT0 + MAX_MINES, min_row=3, max_row=LAST_ROW), titles_from_data=False)
    s = bc_.series[0]
    s.tx = SeriesLabel(strRef=StrRef(f"'Engine'!${E_TOTAL}$2"))
    s.graphicalProperties = GraphicalProperties(solidFill=INK2, ln=LineProperties(noFill=True))
    bc_.set_categories(dates)
    bc_.legend = None
    bc_.x_axis.number_format = "mmm d"
    bc_.x_axis.delete = False
    bc_.y_axis.delete = False
    gen.add_chart(bc_, "A44")


def slim(path: Path):
    """Drop the default theme and docProps and recompress: the file is uploaded as base64."""
    drop = {"xl/theme/theme1.xml", "docProps/app.xml", "docProps/core.xml"}
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist() if n not in drop}
    parts["[Content_Types].xml"] = re.sub(
        rb'<Override PartName="/(xl/theme/theme1|docProps/(app|core))\.xml"[^>]*/>', b"", parts["[Content_Types].xml"])
    parts["_rels/.rels"] = re.sub(rb'<Relationship [^>]*Target="/?docProps/[^"]*"[^>]*/>', b"", parts["_rels/.rels"])
    parts["xl/_rels/workbook.xml.rels"] = re.sub(
        rb'<Relationship [^>]*Target="[^"]*theme1\.xml"[^>]*/>', b"", parts["xl/_rels/workbook.xml.rels"])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, data in parts.items():
            z.writestr(n, data)


def main():
    p = Params()
    wb = Workbook()
    gen = build_generator(wb, p)
    build_engine(wb)
    build_model(wb)
    build_data(wb)
    build_guide(wb)
    add_charts(wb, gen)
    wb.active = 0
    wb.save(OUT)
    slim(OUT)
    print("saved", OUT, OUT.stat().st_size, "bytes")
    if "--lite" in sys.argv:
        save_lite(wb, OUT.with_name("lite.xlsx"))


def save_lite(wb, path: Path):
    """Formulas and named ranges only (no styles, charts, hints, guide): a small file to test the logic."""
    lite = Workbook()
    lite.remove(lite.active)
    for ws in wb.worksheets:
        if ws.title == "Guide":
            continue
        out = lite.create_sheet(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and not (ws.title == "Generator" and c.column == 10):
                    out[c.coordinate] = c.value
    for key, dn in wb.defined_names.items():
        lite.defined_names[key] = DefinedName(key, attr_text=dn.attr_text)
    lite.save(path)
    slim(path)
    print("saved", path, path.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
